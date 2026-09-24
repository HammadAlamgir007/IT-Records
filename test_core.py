"""Self-check for the data layer:  python test_core.py"""

import tempfile
from pathlib import Path

import core
from core import (COLUMNS, COMPANY_FIELDS, Denied, Invalid, Store, export_excel, export_pdf,
                  hash_password, verify_password)


def expect(condition, message):
    if not condition:
        raise AssertionError(message)


def raises(exc, fn, message):
    try:
        fn()
    except exc:
        return
    raise AssertionError(message)


def test_employees(store):
    row_id = store.add_employee(
        {"emp_id": "E-1001", "emp_name": "Bilal Ahmed", "department": "Finance",
         "laptop_serial": "5CG1234ABC", "laptop_name_type": "Dell Latitude 7420",
         "vendor": "Al-Rehman Traders", "dc_no": "DC-991"}, "superadmin", "superadmin")
    row = store.employee(row_id)
    expect(row["emp_name"] == "Bilal Ahmed", "name should be stored")
    expect(row["mobile_serial"] == "", "untouched fields default to empty, never None")
    expect(row["created_by"] == "superadmin", "the creator is recorded")

    # only the Employee ID is compulsory
    store.add_employee({"emp_id": "E-1002"}, "superadmin", "superadmin")
    raises(Invalid, lambda: store.add_employee({"emp_name": "No ID"}, "superadmin", "superadmin"),
           "a missing Employee ID must be refused")
    raises(Invalid, lambda: store.add_employee({"emp_id": "E-1001"}, "superadmin", "superadmin"),
           "a duplicate Employee ID must be refused")

    # search reaches every column, not just the name
    expect(len(store.employees()) == 2, "both records should be listed")
    expect(len(store.employees("bilal")) == 1, "search should ignore case")
    expect(len(store.employees("5CG1234")) == 1, "search should reach the laptop serial")
    expect(len(store.employees("Al-Rehman")) == 1, "search should reach the vendor")
    expect(len(store.employees("nothing here")) == 0, "a miss returns nothing")

    store.update_employee(row_id, {**row, "department": "Audit"}, "superadmin", "superadmin")
    expect(store.employee(row_id)["department"] == "Audit", "the edit should stick")
    expect(store.employee(row_id)["emp_name"] == "Bilal Ahmed", "other fields survive an edit")
    return row_id


def test_rights(store, row_id):
    row = store.employee(row_id)
    # a plain user may look, and nothing else
    raises(Denied, lambda: store.add_employee({"emp_id": "X-1"}, "amir", "user"),
           "a normal user must not add")
    raises(Denied, lambda: store.update_employee(row_id, row, "amir", "user"),
           "a normal user must not edit")
    raises(Denied, lambda: store.delete_employee(row_id, "amir", "user"),
           "a normal user must not delete")
    raises(Denied, lambda: store.logs("user"), "a normal user must not read the logs")
    raises(Denied, lambda: store.users("user"), "a normal user must not manage users")
    expect(len(store.employees()) == 2, "but a normal user's view is not restricted")

    # anything that is not one of the two identities gets nothing at all
    for right in ("view", "edit", "delete", "users", "logs", "export"):
        expect(not Store.may("admin", right), "a removed role keeps no rights")
        expect(not Store.may("", right), "an empty role keeps no rights")

    new_id = store.add_employee({"emp_id": "E-1003", "emp_name": "Sana Khan"},
                                "superadmin", "superadmin")
    store.update_employee(new_id, {"emp_id": "E-1003", "emp_name": "Sana K."},
                          "superadmin", "superadmin")
    expect(store.employee(new_id)["updated_by"] == "superadmin", "the editor is recorded")

    expect(Store.may("user", "view") and not Store.may("user", "edit"), "user views only")
    expect(all(Store.may("superadmin", r) for r in
               ("view", "edit", "delete", "users", "logs", "export")), "super admin has it all")
    expect(set(core.ROLES) == {"user", "superadmin"}, "exactly two identities exist")
    expect(Store.may("viewer", "view") and not Store.may("viewer", "edit"),
           "the unsigned-in viewer can view and nothing else")
    raises(Denied, lambda: store.add_employee({"emp_id": "X-2", "emp_name": "V"},
                                              "anyone", "viewer"),
           "a viewer must not add")
    expect("viewer" not in core.ROLES, "viewer is a session state, never a stored role")


def test_assets(store):
    # add / read / edit across all four registers
    laptop = store.add_asset("laptop", {"name_type": "ThinkPad", "serial": "TP-X1",
                                        "spec": "16GB", "status": "Spare"},
                             "superadmin", "superadmin")
    expect(store.asset(laptop)["identity"] == "TP-X1", "a laptop is keyed by its serial")
    store.update_asset(laptop, {"serial": "TP-X2", "spec": "8GB"},
                       "superadmin", "superadmin")
    expect(store.asset(laptop)["identity"] == "TP-X2", "update follows the new serial")

    def raises_invalid(fn, message):
        try:
            fn()
        except Invalid:
            return
        raise AssertionError(message)

    raises_invalid(lambda: store.add_asset("laptop", {"serial": "TP-X2", "status": "In use"},
                                           "superadmin", "superadmin"),
                   "a duplicate laptop serial must be refused")
    raises_invalid(lambda: store.add_asset("laptop", {}, "superadmin", "superadmin"),
                   "an asset with nothing to identify it must be refused")

    # identity choice per kind
    mob = store.add_asset("mobile", {"imei": "351234567890", "serial": "M-SER",
                                     "sim_number": "0300-1234567"},
                          "superadmin", "superadmin")
    expect(store.asset(mob)["identity"] == "351234567890", "a mobile is keyed by its IMEI")
    ip = store.add_asset("ip_phone", {"mac_no": "AA:BB:CC:00:11:22", "ip_no": "10.0.0.5"},
                         "superadmin", "superadmin")
    expect(store.asset(ip)["identity"] == "AA:BB:CC:00:11:22", "an IP phone by its MAC")
    prn = store.add_asset("printer", {"name_type": "LaserJet M428"}, "superadmin", "superadmin")
    expect(store.asset(prn)["identity"].startswith("PRN-"), "no serial, a tag is issued")

    # assign -> employee sheet syncs
    store.assign_asset(laptop, "E-1001", note="new joiner", actor="superadmin", role="superadmin")
    expect(store.asset(laptop)["current_emp_id"] == "E-1001", "the asset records its holder")
    expect(store.employees("E-1001")[0]["laptop_serial"] == "TP-X2",
           "assigning fills the employee's matching field")
    expect(store.employees("E-1001")[0]["laptop_name_type"] == "ThinkPad",
           "and the model lands too")
    expect(store.asset(laptop)["status"] == "In use",
           "an item somebody holds reads 'In use', not the 'Spare' it was added as")

    # transfer -> the old holder is written off and the new one opens
    store.assign_asset(laptop, "E-1002", note="handover", actor="superadmin", role="superadmin")
    hist = store.asset_history(store.asset(laptop))
    expect(len(hist) == 2, "two lines after an assign + a transfer")
    expect(hist[0]["released_on"] and hist[1]["released_on"] == "",
           "the old line is closed, the new line is open")
    expect([h["emp_id"] for h in hist] == ["E-1001", "E-1002"],
           "history reads oldest first: the first holder, then the current one")
    expect(store.employees("E-1001")[0]["laptop_serial"] == "",
           "the old holder's serial is cleared on the transfer")

    # an employee cannot be given an assets that's already theirs
    raises_invalid(lambda: store.assign_asset(laptop, "E-1002", actor="superadmin", role="superadmin"),
                   "re-assigning to the same holder is refused")

    # a backdated From date lands on the assignment line
    store.assign_asset(laptop, "E-1001", note="handover", actor="superadmin", role="superadmin",
                       start_date="2020-01-02")
    expect(store.asset_history(store.asset(laptop))[-1]["assigned_on"] == "2020-01-02",
           "the chosen From date is recorded")

    # release -> back to the pool, holder still readable in history
    store.release_asset(laptop, note="returned", actor="superadmin", role="superadmin")
    latest = store.asset(laptop)
    expect(latest["current_emp_id"] == "" and latest["current_emp_name"] == "",
           "released removes the current holder")
    expect(latest["status"] == "Spare",
           "and puts it back in stock, so no row says 'In use' with nobody holding it")
    expect(store.asset_history(latest)[0]["released_on"], "the release is time-stamped")
    expect(store.employee_assets("E-1002") and
           store.employee_assets("E-1002")[0]["asset_identity"] == "TP-X2",
           "an employee's asset history lists what they held")

    # the same item may not go to a non-existent employee
    raises_invalid(lambda: store.assign_asset(laptop, "NOPE-1", actor="superadmin", role="superadmin"),
                   "assigning to an unknown employee is refused")

    # rights apply to the registers too
    raises(Denied, lambda: store.add_asset("laptop", {"serial": "V1"}, "v", "viewer"),
           "a viewer must not add assets")
    raises(Denied, lambda: store.assign_asset(laptop, "E-1002", actor="amir", role="user"),
           "a plain user must not assign assets")
    raises(Denied, lambda: store.delete_asset(prn, "amir", "user"),
           "a plain user must not delete assets")
    expect(len(store.assets("laptop")) == 1, "a viewer's view of the registers is not restricted")

    # asset operations leave a trace in the activity log
    expect(len(store.logs("superadmin", "ThinkPad")) > 0, "an asset add/edit is logged")

    # history survives an asset deletion (the row is gone, the log keeps it)
    store.delete_asset(prn, "superadmin", "superadmin")
    expect(store.asset(prn) is None, "the printer is gone")
    expect(len(store.logs("superadmin", "LaserJet")) > 0, "but its deletion is logged")


def test_asset_seeding(folder):
    """Serial numbers already sitting on employee records become asset rows with
    history the first time the file is opened after the upgrade."""
    path = folder / "seed.db"
    store = Store(path)
    store.bootstrap()
    store.add_employee({"emp_id": "E-3001", "emp_name": "Seedy", "laptop_serial": "SEED-LAP",
                        "laptop_name_type": "ThinkPad", "printer_serial": "SEED-PRN",
                        "mac_no": "BB:CC:DD:00:11:22"},
                       "superadmin", "superadmin")

    reopened = Store(path)                     # opening again runs the one-time migration
    expect([a["identity"] for a in reopened.assets("laptop")] == ["SEED-LAP"],
           "the laptop serial is seeded as an asset")
    expect([a["identity"] for a in reopened.assets("printer")] == ["SEED-PRN"],
           "the printer serial is seeded as an asset")
    expect(len(reopened.assets("ip_phone")) == 1, "a MAC number becomes an IP phone asset")
    seeded = reopened.assets("laptop")[0]
    expect(seeded["current_emp_id"] == "E-3001", "the seeded holder is the employee")
    expect(reopened.asset_history(seeded)[0]["emp_id"] == "E-3001",
           "and the history is written from day one")

    # opening it again does not duplicate anything
    again = Store(path)
    expect(len(again.assets("laptop")) == 1, "a catch-up run never duplicates")


def test_logs(store, row_id):
    store.delete_employee(row_id, "superadmin", "superadmin")
    expect(store.employee(row_id) is None, "the record should be gone")

    logs = store.logs("superadmin")
    expect(logs[0]["action"] == "delete", "the newest entry comes first")
    expect("Bilal Ahmed" in logs[0]["detail"], "a deletion keeps what was deleted")
    expect(logs[0]["actor"] == "superadmin", "the log records who did it")
    expect(logs[0]["at"].startswith("20"), "the log records when")

    edits = store.logs("superadmin", "Department")
    expect(any("'Finance' -> 'Audit'" in e["detail"] for e in edits),
           "an edit records the old and the new value")
    expect(len(store.logs("superadmin", "no-such-thing")) == 0, "log search can miss")


def test_accounts(store):
    expect(store.bootstrap() is None, "bootstrap must not run twice")
    store.create_user("amir", "view-only-pw", "user", "superadmin", "superadmin")
    raises(Denied, lambda: store.create_user("x", "password1", "user", "amir", "user"),
           "a normal user must not create users")
    raises(Invalid, lambda: store.create_user("bob", "short", "user", "superadmin", "superadmin"),
           "a short password must be refused")
    raises(Invalid, lambda: store.create_user("zoe", "password1", "boss", "superadmin", "superadmin"),
           "an unknown role must be refused")
    raises(Invalid, lambda: store.create_user("amir", "password1", "user", "superadmin", "superadmin"),
           "a duplicate username must be refused")

    expect(store.login("amir", "view-only-pw") == "user", "a correct password signs in")
    expect(store.login("AMIR", "view-only-pw") == "user", "the username is not case sensitive")
    raises(Invalid, lambda: store.login("amir", "wrong"), "a wrong password is refused")
    raises(Invalid, lambda: store.login("ghost", "whatever"), "an unknown user is refused")

    # own password yes, someone else's no
    store.set_password("amir", "new-password", "amir", "user")
    expect(store.login("amir", "new-password") == "user", "the new password works")
    raises(Denied, lambda: store.set_password("superadmin", "hijacked!", "amir", "user"),
           "a user must not change another user's password")

    raises(Invalid, lambda: store.delete_user("superadmin", "superadmin", "superadmin"),
           "you cannot delete your own account")
    store.create_user("second", "password-two", "superadmin", "superadmin", "superadmin")
    store.delete_user("second", "superadmin", "superadmin")

    # five wrong tries and the account is locked, even with the right password
    for _ in range(5):
        raises(Invalid, lambda: store.login("amir", "wrong"), "wrong password")
    raises(Invalid, lambda: store.login("amir", "new-password"), "lockout must bite")

    expect(verify_password("secret-password", hash_password("secret-password")), "hash round trip")
    expect(not verify_password("wrong", hash_password("secret-password")), "wrong password fails")
    expect(hash_password("same") != hash_password("same"), "the hash is salted")
    expect(not verify_password("x", "not-a-hash"), "a corrupt hash fails instead of crashing")


def test_legacy_roles(folder):
    """A database from the three-role version must not leave anyone half-privileged."""
    path = folder / "legacy.db"
    store = Store(path)
    store.bootstrap()
    store.create_user("sara", "some-password", "user", "superadmin", "superadmin")
    import sqlite3
    with sqlite3.connect(path) as c:                     # pretend sara was an old "admin"
        c.execute("UPDATE users SET role='admin' WHERE username='sara'")

    reopened = Store(path)                               # opening the file does the cleanup
    roles = {u["username"]: u["role"] for u in reopened.users("superadmin")}
    expect(roles["sara"] == "user", "an unknown role is reset to view-only")
    expect(roles["superadmin"] == "superadmin", "a valid role is left alone")
    expect(any("unknown role reset" in e["detail"] for e in reopened.logs("superadmin")),
           "and the reset is written to the log")


def test_no_way_in(folder):
    """A database with users but no super admin must hand back a way in, not lock everyone out."""
    path = folder / "locked.db"
    store = Store(path)
    store.bootstrap()
    store.create_user("amir", "view-only-pw", "user", "superadmin", "superadmin")
    import sqlite3
    with sqlite3.connect(path) as c:                 # the only super admin gets demoted
        c.execute("UPDATE users SET role='user' WHERE username='superadmin'")

    reopened = Store(path)
    made = reopened.bootstrap()
    expect(made is not None, "a database with no super admin must make one")
    user, password = made
    expect(reopened.login(user, password) == "superadmin", "and the new password must work")
    expect(reopened.login("amir", "view-only-pw") == "user", "other accounts are left alone")
    expect(reopened.bootstrap() is None, "once there is a super admin it stops interfering")


def test_import(folder):
    """Import a sheet written the way the IT department already writes them."""
    from datetime import datetime

    from openpyxl import Workbook

    from core import match_column, read_excel

    # headings straight off the existing laptop sheet, including ones we ignore
    book = Workbook()
    sheet = book.active
    sheet.append(["Sno", "Emp ID", "Name", "IP Phone", "Department", "User ID", "Host Name",
                  "Email Aik Address", "Contact No", "Remarks", "Location", "Region",
                  "Laptop Serial Number", "Name & Type", "Delivery Dates", "Cost Centre"])
    sheet.append([1, "E-2001", "Imran Yousaf", 2201, "Finance", "imran.y", "PK-HO-101",
                  "imran@company.com.pk", "0300-1234567", "new joiner", "Head Office",
                  "Punjab", "5CG9999XYZ", "Dell Latitude 7420", datetime(2026, 3, 2), "CC-9"])
    sheet.append([2, "E-2002", "Hina Raza", None, "HR", None, None, None, None, None,
                  "Lahore Branch", "Punjab", "5CG8888ABC", "HP ProBook 450", None, None])
    sheet.append([3, None, "No ID Here", None, "IT", None, None, None, None, None,
                  None, None, None, None, None, None])
    sheet.append([None] * 16)                                   # a blank separator row
    sheet.append([5, "E-2002", "Hina R.", None, "HR", None, None, None, None, None,
                  None, None, None, None, None, None])          # same ID twice
    path = folder / "laptops.xlsx"
    book.save(path)

    expect(match_column("Emp ID") == "emp_id", "a known heading maps to its field")
    expect(match_column("Email Aik Address") == "email", "the sheet's own wording is understood")
    expect(match_column("Laptop Serial Number") == "laptop_serial", "serial maps across")
    expect(match_column("Cost Centre") is None, "an unknown heading maps to nothing")

    rows, ignored, problems = read_excel(path)
    expect(len(rows) == 4, "four usable rows: two employees, the no-id row, and the duplicate")
    expect("Cost Centre" in ignored, "unknown headings are reported")
    expect(any("no Employee ID" in p for p in problems),
           "a row without an ID is called out, even though it is still imported")
    expect(any("also on row" in p for p in problems), "a repeated ID is reported")
    expect(rows[0]["ip_phone"] == "2201", "a number becomes text, not 2201.0")
    expect(rows[0]["delivery_date"] == "2026-03-02", "a date becomes a plain date")
    expect(rows[0]["laptop_name_type"] == "Dell Latitude 7420", "Name & Type is the laptop")
    expect(rows[1]["user_id"] == "", "an empty cell becomes empty text, never None")
    expect(rows[2]["emp_id"].startswith("UNASSIGNED"),
           "a missing ID gets a placeholder instead of being dropped")

    store = Store(folder / "import.db")
    store.bootstrap()
    raises(Denied, lambda: store.import_rows(rows, "amir", "user"),
           "a normal user must not import")

    result = store.import_rows(rows, "superadmin", "superadmin")
    expect(result["added"] == 3, "three new records: two employees plus the no-id row")
    expect(result["updated"] == 1, "the repeated ID updates the one just added")
    expect(len(store.employees()) == 3, "and three records exist")
    expect(store.employees("E-2002")[0]["emp_name"] == "Hina R.", "the later row wins")
    expect(store.employees("Imran")[0]["department"] == "Finance", "fields land in the right place")

    # importing the same file again must not duplicate anyone
    again = store.import_rows(rows, "superadmin", "superadmin")
    expect(again["added"] == 0 and len(store.employees()) == 3, "a second import adds nobody")

    # a sheet that carries only some columns must not wipe the rest
    thin = [{"emp_id": "E-2001", "department": "Audit", "_row": 2}]
    store.import_rows(thin, "superadmin", "superadmin")
    updated = store.employees("E-2001")[0]
    expect(updated["department"] == "Audit", "the supplied field is updated")
    expect(updated["emp_name"] == "Imran Yousaf", "and the untouched fields survive")

    # the 'leave existing alone' choice
    skipped = store.import_rows(thin, "superadmin", "superadmin", update_existing=False)
    expect(skipped["skipped"] == 1 and skipped["updated"] == 0, "existing rows can be left alone")

    # an Employee ID that already belongs to someone else must not be renamed over
    renamed = [{"emp_id": "E-2001", "emp_name": "Someone Else", "_row": 2}]
    conflict = store.import_rows(renamed, "superadmin", "superadmin")
    expect(conflict["skipped"] == 1 and conflict["updated"] == 0,
           "a name mismatch on an existing ID is skipped, not overwritten")
    expect(store.employees("E-2001")[0]["emp_name"] == "Imran Yousaf",
           "the real Imran Yousaf is untouched by the mismatched row")

    expect(any(e["action"] == "import" for e in store.logs("superadmin")),
           "the import itself is logged")

    # a file that is not a sheet of employees at all
    other = Workbook()
    other.active.append(["Fruit", "Colour"])
    other.active.append(["Mango", "Yellow"])
    other.save(folder / "fruit.xlsx")
    rows2, _, problems2 = read_excel(folder / "fruit.xlsx")
    expect(rows2 == [] and problems2, "an unrelated sheet is refused with a reason")

    # the file our own export writes must import straight back
    export_excel(store.employees(), folder / "round.xlsx")
    back, _, trouble = read_excel(folder / "round.xlsx")
    expect(len(back) == 3 and not trouble, "an exported file re-imports cleanly")


def test_bulk_import(folder):
    """The single-transaction importer must behave exactly like the old per-row one."""
    store = Store(folder / "bulk.db")
    store.bootstrap()
    store.add_employee({"emp_id": "B-1", "emp_name": "Base Person", "department": "Ops"},
                       "superadmin", "superadmin")

    raises(Denied, lambda: store.bulk_import([], "amir", "user"), "a normal user must not import")

    batch = [
        {"emp_id": "B-2", "emp_name": "New One", "department": "IT", "laptop_serial": "BULK-SER",
         "_row": 100},
        {"emp_id": "B-3", "emp_name": "New Two", "_row": 101},
        {"emp_id": "B-2", "emp_name": "New One", "email": "n@x.com", "_row": 102},
        {"emp_id": "B-4", "emp_name": "Sneaky", "laptop_serial": "BULK-SER", "_row": 103},
    ]
    result = store.bulk_import(batch, actor="superadmin", role="superadmin")
    expect(result["added"] == 2 and result["updated"] == 1 and result["skipped"] == 1,
           "one transaction adds two, updates the repeated ID, and refuses the stolen serial")
    expect(store.employees("B-2")[0]["email"] == "n@x.com", "the later row updated the first")
    expect(store.employees("B-3")[0]["department"] == "", "a thin row stays thin")
    expect(not store.employees("B-4"), "the duplicate-serial row was not imported")
    expect(any("BULK-SER" in p for p in result["problems"]), "and the skip said why")

    again = store.bulk_import(batch, actor="superadmin", role="superadmin")
    expect(again["added"] == 0 and again["updated"] == 3 and again["skipped"] == 1,
           "re-running the same batch adds nobody")

    untouched = store.bulk_import(batch, update_existing=False, actor="superadmin", role="superadmin")
    expect(untouched["added"] == 0 and untouched["updated"] == 0 and untouched["skipped"] == 4,
           "existing rows can be left alone in bulk too")

    mismatch = [{"emp_id": "B-2", "emp_name": "Impostor", "_row": 9}]
    miss = store.bulk_import(mismatch, actor="superadmin", role="superadmin")
    expect(miss["skipped"] == 1 and store.employees("B-2")[0]["emp_name"] == "New One",
           "a name mismatch is skipped rather than renamed over")

    # a sheet with no Employee ID at all is refused per row, not for everyone
    no_id = [{"emp_name": "Nobody"}, {"emp_id": "B-5", "emp_name": "Somebody", "_row": 11}]
    res = store.bulk_import(no_id, actor="superadmin", role="superadmin")
    expect(res["added"] == 1 and res["skipped"] == 1 and store.employees("B-5"),
           "only the row that actually has an ID is imported")


def test_source_files(folder):
    """The live three-file import path: read UPDATE / CHALLAN / mobile once, then
    write everything into the database in one batch."""
    here = Path(__file__).parent
    missing = [f for f in ("UPDATE.xlsx", "CHALLAN SHEET.xlsx", "mobile.xls")
               if not (here / f).exists()]
    if missing:
        print(f"  (skipping source-file check: missing {', '.join(missing)})")
        return

    import build_import_file as bif

    try:
        clean, review, needs_id, stats = bif.records_for_import()
    except ModuleNotFoundError as exc:
        # mobile.xls is the old binary Excel format, which needs xlrd. It is a
        # build-time helper, not something the app itself ever reads.
        print(f"  (skipping source-file check: {exc})")
        return
    expect(stats["clean"] > 0 and stats["matched"] > 0,
           "the three files yield rows and challan matches")
    expect(stats["printers"] > 0 and stats["ip"] > 0 and stats["stock"] > 0,
           "and every register is represented")

    store = Store(folder / "src.db")
    store.bootstrap()
    result = store.bulk_import(bif.to_store_records(clean), actor="system", role="superadmin")
    expect(result["added"] == stats["clean"] and result["skipped"] == 0,
           f"all {stats['clean']} source rows import cleanly in one pass")

    store = Store(folder / "src.db")      # reopen: seeding builds the laptop register
    asset_result = store.upsert_assets(bif.all_register_assets(), actor="system", role="superadmin")
    expect(asset_result["added"] + asset_result["updated"] > 0
           and asset_result["skipped"] == 0,
           "the register pass adds stock and enriches owned laptops without skips")

    rows = store.employees()
    expect(len(rows) == stats["clean"],
           "employees are exactly the Laptop sheet's people - no synthetic rows")
    expect(any((r.get("vendor") or "").strip() for r in rows if (r.get("laptop_serial") or "").strip()),
           "challan vendor info lands on the matched laptop rows")
    expect(any(r["emp_id"].startswith("NOID-") for r in rows),
           "people without an ID in the source keep a reviewable placeholder")
    laptops = store.assets("laptop")
    expect(any((r.get("current_emp_name") or "").strip() for r in laptops),
           "owned laptops have their employee as current owner")

    expect(store.assets("printer") and store.assets("ip_phone")
           and store.assets("mobile"),
           "printers, IP phones and mobile stock arrive as register assets")
    expect(not store.employees("PRINTER-001") and not store.employees("IP-001")
           and not store.employees("STOCK-001"),
           "the asset sheets do not create employee rows")
    for kind in ("printer", "ip_phone", "mobile"):
        expect(all(not (r.get("current_emp_name") or "").strip()
                   for r in store.assets(kind)),
               "an asset with no known holder keeps Current Owner empty")

    challan_rows = store.assets("challan")
    expect(len(challan_rows) > 0, "every challan delivery becomes a challan-register row")
    linked = store.link_challan_from_employees(actor="system", role="superadmin")
    expect(linked == stats["matched"],
           "the matched deliveries are marked issued to their employee")
    expect(any((r.get("current_emp_name") or "").strip()
               for r in store.assets("challan")),
           "an issued challan line carries the employee's name")
    expect(any(not (r.get("current_emp_name") or "").strip()
               and r["status"] == "Spare" for r in store.assets("challan")),
           "an unmatched delivery stays spare with no owner")


def test_challan(store):
    """The challan register: a delivery line per item, issue books the machine
    into the Laptops register, and releasing that laptop writes the line back
    to stock - the two registers agree."""
    def raises_invalid(fn, message):
        try:
            fn()
        except Invalid:
            return
        raise AssertionError(message)

    ch = store.add_asset("challan", {"serial": "CHL-SER-9", "vendor": "Acer House",
                                     "name_type": "Laptop 15", "spec": "Core i5",
                                     "dc_no": "DC-9", "delivery_date": "2026-01-05",
                                     "notes": "office order"}, "superadmin", "superadmin")
    expect(store.asset(ch)["identity"] == "CHL-SER-9",
           "a challan line is keyed by the item's serial")
    expect(store.asset(ch)["status"] == "In use",
           "the challan register uses the same status choices")

    summary = store.issue_challan(ch, "E-1001", note="delivery to new joiner",
                                  actor="superadmin", role="superadmin")
    laps = [a for a in store.assets("laptop") if a["serial"] == "CHL-SER-9"]
    expect(len(laps) == 1, "issuing a challan line books the machine into Laptops")
    expect(laps[0]["current_emp_id"] == "E-1001", "and the machine is issued to the employee")
    expect(laps[0]["vendor"] == "Acer House" and laps[0]["dc_no"] == "DC-9",
           "the laptop inherits the delivery's vendor and DC")
    challan = store.asset(ch)
    expect(challan["current_emp_id"] == "E-1001" and challan["status"] == "In use",
           "the challan line records the same holder")
    expect(store.employees("E-1001")[0]["laptop_serial"] == "CHL-SER-9",
           "the employee's matching field is filled from the delivery")
    expect(store.employee_assets("E-1001")[0]["asset_identity"] == "CHL-SER-9",
           "the delivery shows up in the employee's asset history")
    expect("handed" not in summary and "assigned" in summary,
           "a fresh issue reads as an assignment")

    raises_invalid(lambda: store.issue_challan(ch, "E-1001",
                                               actor="superadmin", role="superadmin"),
                   "re-issuing the same line to the same holder is refused")
    store.issue_challan(ch, "E-1002", note="handover", actor="superadmin", role="superadmin")
    expect(store.asset(ch)["current_emp_id"] == "E-1002",
           "re-issuing a challan line moves it to the new holder")

    store.release_asset(laps[0]["id"], note="returned", actor="superadmin", role="superadmin")
    expect(store.asset(ch)["current_emp_id"] == "" and store.asset(ch)["status"] == "Spare",
           "releasing the laptop writes its challan line back to stock")

    raises(Denied, lambda: store.add_asset("challan", {"serial": "X1"}, "v", "viewer"),
           "a viewer must not add challan lines")
    raises_invalid(lambda: store.issue_challan(ch, "NOPE-1",
                                               actor="superadmin", role="superadmin"),
                   "issuing to an unknown employee is refused")


def test_settings(store):
    def raises_invalid(fn, message):
        try:
            fn()
        except Invalid:
            return
        raise AssertionError(message)

    detail = store.company()
    expect(set(detail) == set(COMPANY_FIELDS) and not detail["name"],
           "a fresh database has empty organisation details")

    raises(Denied, lambda: store.save_company({"name": "ABC Corp"}, "v", "viewer"),
           "a viewer must not edit the organisation detail")
    raises_invalid(lambda: store.save_company({"name": "  "}, "s", "superadmin"),
                   "a company name is required for the letterhead")

    saved = store.save_company({"name": "ABC Corporation (Pvt) Ltd", "address": "Main Road",
                                "city": "Islamabad", "phone": "051-1234567",
                                "email": "it@abc.com"}, "superadmin", "superadmin")
    expect(saved["name"] == "ABC Corporation (Pvt) Ltd", "the cleaned details are returned")
    expect(store.company()["city"] == "Islamabad" and store.company()["name"] == saved["name"],
           "the letterhead settings stick")
    expect(store.logs("superadmin", "organisation")[0]["action"] == "setting",
           "editing the organisation detail is audited")


def test_record_pdf(store, folder):
    from core import export_record_pdf
    company = store.company()
    history = [{"emp": "Bilal Ahmed (E-1001)", "from": "05 Jan 2026", "until": "",
                "note": "first issue"}]
    pdf = export_record_pdf("LAPTOP RECORD",
                            [("Serial No", "5CD1234ABC"), ("Name & Type", "HP PRO BOOK")],
                            company=company, history=history, document_no="ITR/LAP/5CD1234ABC",
                            prepared_by="superadmin",
                            path=folder / "record.pdf")
    expect(pdf.exists() and pdf.read_bytes().startswith(b"%PDF") and pdf.stat().st_size > 0,
           "a formal one-record PDF is written")

    empty = export_record_pdf("ASSET RECORD", [], company=company, path=folder / "empty.pdf")
    expect(empty.stat().st_size > 0, "a record with no details still prints")

    check_pdf_type_size(pdf, "the record document")
    check_pdf_type_size(export_pdf(store.employees(), folder / "table.pdf"), "the register table")


def check_pdf_type_size(path, what):
    """Setting a page size on a QTextDocument makes Qt scale the finished layout
    onto the paper: every point size comes out far bigger than the stylesheet
    asked for and the right-hand side runs off the page. Read the type back out
    of the file and make sure it is the size that was asked for."""
    try:
        import pymupdf
    except ImportError:
        print("  (skipping PDF type-size check: pymupdf is not installed)")
        return
    with pymupdf.open(path) as doc:
        page = doc[0]
        spans = [s for b in page.get_text("dict")["blocks"]
                 for l in b.get("lines", []) for s in l["spans"]]
        sizes = [s["size"] for s in spans]
        right = max((s["bbox"][2] for s in spans), default=0)
        width = page.rect.width
    expect(sizes and max(sizes) <= 30,
           f"{what}: nothing is printed at a runaway size (largest was {max(sizes):.1f}pt)")
    expect(right <= width,
           f"{what}: no text runs off the right-hand edge of the paper")


def test_exports(store, folder):
    rows = store.employees()
    xlsx = export_excel(rows, folder / "employees.xlsx")
    expect(xlsx.exists() and xlsx.stat().st_size > 0, "the Excel file should be written")

    from openpyxl import load_workbook
    sheet = load_workbook(xlsx).active
    expect(sheet.max_row == len(rows) + 1, "one header row plus one row per employee")
    expect(sheet.cell(1, 2).value == "Employee ID", "the header uses the screen labels")
    expect(sheet.max_column == len(COLUMNS), "every field, and no extra counter column")
    headings = [c.value for c in sheet[1]]
    expect(len([h for h in headings if str(h).strip().lower() == "sno"]) == 1,
           "exactly one Sno column - the sheet's own, not a duplicate counter")

    pdf = export_pdf(rows, folder / "employees.pdf", subtitle="Self-test")
    expect(pdf.exists() and pdf.stat().st_size > 0, "the PDF should be written")
    expect(pdf.read_bytes().startswith(b"%PDF"), "and it should really be a PDF")

    # a value Excel would otherwise run as a formula
    export_excel([{**rows[0], "remarks": "=cmd()"}], folder / "risky.xlsx")
    risky = load_workbook(folder / "risky.xlsx").active
    values = [c.value for c in risky[2]]
    remarks_cell = values[COLUMNS.index("remarks")]
    expect(str(remarks_cell).startswith("'="), "a formula-looking value is neutralised")
    blank_cell = values[COLUMNS.index("designation")]       # left blank in this row
    expect(blank_cell in (None, ""), "a blank field stays blank, not a stray quote mark")

    # The employee sheet has 30-odd fields and most are blank; printing them all
    # squeezed every value into a sliver a letter wide, so the PDF drops the
    # columns nothing fills.
    sparse = [{**{c: "" for c in COLUMNS}, "emp_id": "E-1", "emp_name": "Only Two"}]
    kept = core._columns_with_data(sparse, COLUMNS)
    expect(kept == ["emp_id", "emp_name"], "an all-blank column is left off the printed page")
    expect(core._columns_with_data([{c: "" for c in COLUMNS}], COLUMNS) == COLUMNS,
           "but a page with nothing on it at all still prints its headings")


def test_fixed_password(folder: Path):
    """The super admin password is a constant, so a database made by an older
    build - or one whose password was changed - is put back on it when opened."""
    path = folder / "fixed.db"
    store = Store(path)
    store.bootstrap()
    expect(store.login(core.SUPERADMIN_USER, core.SUPERADMIN_PASSWORD) == "superadmin",
           "the fixed password signs in")

    store.reset_superadmin("something-else-entirely")
    raises(Invalid, lambda: store.login(core.SUPERADMIN_USER, core.SUPERADMIN_PASSWORD),
           "a changed password really does take effect")

    reopened = Store(path)
    expect(reopened.bootstrap() is None, "reopening reports no new account")
    expect(reopened.login(core.SUPERADMIN_USER, core.SUPERADMIN_PASSWORD) == "superadmin",
           "and puts the fixed password back, so the login is never lost")


def test_sno(folder: Path):
    """SNO is a gap-free 1, 2, 3 ... sequence the app owns: new and imported
    rows join the end, a delete closes the gap, and the register lists in it."""
    path = folder / "sno.db"
    store = Store(path)
    store.bootstrap()
    blank = {c: "" for c in COLUMNS}
    store.add_employee(blank | {"emp_id": "S-1", "sno": "7"}, "system", "superadmin")
    store.add_employee(blank | {"emp_id": "S-2"}, "system", "superadmin")
    by_id = {r["emp_id"]: r for r in store.employees()}
    expect(by_id["S-1"]["sno"] == "1", "numbering starts at 1 whatever was typed")
    expect(by_id["S-2"]["sno"] == "2", "and a new employee joins the end")

    # An import (its own Sno column is ignored for new rows - they join the end).
    store.bulk_import([{"emp_id": "S-3", "emp_name": "Imported", "sno": "1", "_row": 2}],
                      actor="system", role="superadmin")
    expect({r["emp_id"]: r for r in store.employees()}["S-3"]["sno"] == "3",
           "an imported row is numbered on the end too")

    # Editing never moves a row.
    s1 = {r["emp_id"]: r for r in store.employees()}["S-1"]
    store.update_employee(s1["id"], s1 | {"sno": "99", "emp_name": "Renamed"},
                          "system", "superadmin")
    expect({r["emp_id"]: r for r in store.employees()}["S-1"]["sno"] == "1",
           "an edit keeps the row's place")

    # A delete closes the gap it leaves.
    s2 = {r["emp_id"]: r for r in store.employees()}["S-2"]
    store.delete_employee(s2["id"], "system", "superadmin")
    listed = store.employees()
    expect([r["sno"] for r in listed] == ["1", "2"], "after a delete the numbers stay 1..N")
    expect([r["emp_id"] for r in listed] == ["S-1", "S-3"], "and the register lists in SNO order")

    # An old database: the column was added by a migration, so every row was left
    # blank. Opening the file fills them in.
    with store._conn() as c:
        c.execute("UPDATE employees SET sno=''")
    filled = Store(path).employees()
    expect(all(r["sno"].strip() for r in filled), "reopening fills in every missing number")
    expect(len({r["sno"] for r in filled}) == len(filled), "and no two rows share one")


def test_bundle_seed(folder: Path):
    """A real database shipped beside the program (YOUR-DATA folder) seeds a
    fresh computer's profile on first open, so a bare exe click is not empty."""
    import core
    untouched = folder / "kept.db"
    core.Store(untouched).bootstrap()   # this PC already has data, made first...
    core.Store(untouched).add_employee({c: "" for c in core.COLUMNS} | {"emp_id": "D-1",
                                                                        "emp_name": "Don't touch"}, "system", "superadmin")

    live = folder / "live.db"
    store = core.Store(live)
    store.bootstrap()
    store.add_employee({c: "" for c in core.COLUMNS} | {"emp_id": "RB-1",
                                                        "emp_name": "Bundled"}, "system", "superadmin")

    shipped = folder / "box" / "YOUR-DATA"
    shipped.mkdir(parents=True)
    (shipped / "employees.db").write_bytes(live.read_bytes())

    original = core._bundled_db
    core._bundled_db = lambda: shipped / "employees.db"
    try:
        blank = folder / "empty-profile" / "employees.db"
        seeded = core.Store(blank)
        expect(seeded.employees() and any(e["emp_name"] == "Bundled" for e in seeded.employees()),
               "a fresh profile is seeded from the bundled database")
        again = core.Store(untouched)
        expect(len(again.employees()) == 1 and again.employees()[0]["emp_name"] == "Don't touch",
               "an existing database is never overwritten by the bundle")
    finally:
        core._bundled_db = original


def main():
    with tempfile.TemporaryDirectory() as tmp:
      try:
        folder = Path(tmp)
        store = Store(folder / "test.db")
        user, password = store.bootstrap()
        expect(user == core.SUPERADMIN_USER and password == core.SUPERADMIN_PASSWORD,
               "first run creates the super admin on the fixed password")
        expect(store.login(user, password) == "superadmin", "the printed password works")
        test_fixed_password(folder)
        test_sno(folder)

        row_id = test_employees(store)
        test_rights(store, row_id)
        test_assets(store)
        test_challan(store)
        test_settings(store)
        test_record_pdf(store, folder)
        test_logs(store, row_id)
        test_accounts(store)
        test_legacy_roles(folder)
        test_no_way_in(folder)
        test_asset_seeding(folder)
        test_bulk_import(folder)
        test_source_files(folder)
        test_exports(store, folder)
        test_import(folder)
        test_bundle_seed(folder)
      finally:
        # Close the log file whatever happened, or Windows refuses to delete the
        # temp folder and the real failure is buried under a PermissionError.
        import logging
        logging.shutdown()

    print("self-check OK")


if __name__ == "__main__":
    main()
